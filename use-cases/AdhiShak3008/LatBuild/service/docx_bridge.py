"""
docx_bridge.py
Convert PDF-derived section prose into a contextual DOCX for SuperDocs.

The whole paper is sent to SuperDocs for context, but selected sections are
wrapped in explicit LatBuild boundary markers. Those markers are temporary
transport metadata and are never written back to the PDF.
"""

import io
import re
from docx import Document
from docx.shared import Pt
from html.parser import HTMLParser


def prose_to_docx(text: str, title: str = "Section") -> bytes:
    """Convert plain text to a minimal .docx file."""
    doc = Document()
    for para in list(doc.paragraphs):
        para._element.getparent().remove(para._element)

    paragraphs = [p.strip() for p in text.split("\n\n") if p.strip()]
    if not paragraphs:
        paragraphs = [text.strip()] if text.strip() else ["(empty)"]

    for para_text in paragraphs:
        p = doc.add_paragraph(para_text)
        p.style.font.size = Pt(11)

    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


def _marker_id(title: str) -> str:
    """Stable readable marker token for a section title."""
    token = re.sub(r"[^A-Za-z0-9]+", "_", title).strip("_").upper()
    return token or "SECTION"


def full_paper_to_docx(fm, target_section_titles: list) -> bytes:
    """
    Build the entire paper for SuperDocs.

    Every selected section receives explicit START/END transport markers.
    SuperDocs can therefore read the whole paper while LatBuild retains an
    unambiguous semantic boundary around every authorized write target.
    """
    targets = set(target_section_titles or [])
    doc = Document()

    for para in list(doc.paragraphs):
        para._element.getparent().remove(para._element)

    for section in fm.sections:
        if not section.full_text and not section.heading_block:
            continue

        selected = section.title in targets
        marker = _marker_id(section.title)

        if selected:
            start = doc.add_paragraph(
                f"[[LATBUILD_SECTION_START:{marker}]]"
            )
            start.style.font.size = Pt(7)

        h = doc.add_heading(section.title, level=section.level)
        h.style.font.size = Pt(11)

        if section.full_text:
            paragraphs = [
                p.strip()
                for p in section.full_text.split("\n\n")
                if p.strip()
            ]
            for para_text in paragraphs:
                p = doc.add_paragraph(para_text)
                p.style.font.size = Pt(10)

        if selected:
            end = doc.add_paragraph(
                f"[[LATBUILD_SECTION_END:{marker}]]"
            )
            end.style.font.size = Pt(7)

    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


class _HTMLTextExtractor(HTMLParser):
    """Strip HTML tags and return plain text."""
    def __init__(self):
        super().__init__()
        self._parts = []
        self._skip_tags = {"script", "style"}
        self._in_skip = False

    def handle_starttag(self, tag, attrs):
        if tag in self._skip_tags:
            self._in_skip = True
        if tag in ("p", "div", "br", "h1", "h2", "h3", "h4", "h5", "h6", "li"):
            self._parts.append("\n")

    def handle_endtag(self, tag):
        if tag in self._skip_tags:
            self._in_skip = False

    def handle_data(self, data):
        if not self._in_skip:
            self._parts.append(data)

    def get_text(self):
        return "".join(self._parts)


def extract_text_from_docx(docx_bytes: bytes) -> str:
    """Extract plain text from .docx bytes."""
    buf = io.BytesIO(docx_bytes)
    doc = Document(buf)
    paragraphs = [para.text.strip() for para in doc.paragraphs if para.text.strip()]
    return "\n\n".join(paragraphs)


def extract_text_from_html(html: str) -> str:
    """Extract plain text from SuperDocs HTML."""
    extractor = _HTMLTextExtractor()
    extractor.feed(html)
    raw = extractor.get_text()

    lines = [line.strip() for line in raw.split("\n")]
    cleaned = []
    blank_count = 0
    for line in lines:
        if not line:
            blank_count += 1
        else:
            if blank_count > 0:
                cleaned.append("")
            blank_count = 0
            cleaned.append(line)

    return "\n\n".join(
        para for para in "\n".join(cleaned).split("\n\n") if para.strip()
    )


def extract_marked_section_from_text(
    full_text: str, section_title: str
) -> str | None:
    """
    Extract the text between LatBuild START/END markers.

    Returns None when the markers are absent or malformed. We intentionally
    fail closed instead of guessing a section boundary.
    """
    marker = _marker_id(section_title)
    pattern = re.compile(
        rf"\[\[LATBUILD_SECTION_START:{re.escape(marker)}\]\]\s*"
        rf"(.*?)"
        rf"\[\[LATBUILD_SECTION_END:{re.escape(marker)}\]\]",
        re.DOTALL | re.IGNORECASE,
    )
    match = pattern.search(full_text or "")
    if not match:
        return None

    body = match.group(1).strip()

    # Remove the heading itself; the heading is not editable body text.
    heading = section_title.strip()
    if body.lower().startswith(heading.lower()):
        body = body[len(heading):].lstrip()

    return body.strip() or ""
